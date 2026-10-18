from __future__ import annotations

import base64
import tempfile
import unittest
from io import BytesIO
from pathlib import Path
from uuid import uuid4

from openpyxl import Workbook, load_workbook

from membership import (
    WORKBOOK_FIELDS,
    decrypt_record,
    decrypt_record_fields,
    decode_encryption_key,
    encrypt_record,
    encrypt_record_fields,
    membership_workbook,
    new_membership_records,
    search_records,
    workbook_records,
)


class EncryptionTests(unittest.TestCase):
    def setUp(self) -> None:
        self.key = bytes(range(32))
        self.record_id = str(uuid4())
        self.record = {"FIRST NAME": "Ama", "DATE OF BIRTH": "2000-02-03"}

    def test_round_trip_uses_versioned_authenticated_envelope(self) -> None:
        encrypted = encrypt_record(self.record, self.record_id, self.key)
        self.assertTrue(encrypted.startswith("v1:"))
        self.assertNotIn("Ama", encrypted)
        self.assertEqual(decrypt_record(encrypted, self.record_id, self.key), self.record)

    def test_tampered_envelope_is_rejected(self) -> None:
        encrypted = encrypt_record(self.record, self.record_id, self.key)
        prefix, token = encrypted.split(":", 1)
        replacement = "A" if token[0] != "A" else "B"
        tampered = f"{prefix}:{replacement}{token[1:]}"
        with self.assertRaisesRegex(ValueError, "could not be authenticated"):
            decrypt_record(tampered, self.record_id, self.key)

    def test_ciphertext_cannot_be_moved_to_another_record(self) -> None:
        encrypted = encrypt_record(self.record, self.record_id, self.key)
        with self.assertRaisesRegex(ValueError, "could not be authenticated"):
            decrypt_record(encrypted, str(uuid4()), self.key)

    def test_wrong_key_is_rejected(self) -> None:
        encrypted = encrypt_record(self.record, self.record_id, self.key)
        with self.assertRaisesRegex(ValueError, "could not be authenticated"):
            decrypt_record(encrypted, self.record_id, bytes(reversed(range(32))))

    def test_key_must_be_valid_base64_and_32_bytes(self) -> None:
        encoded = base64.b64encode(self.key).decode()
        self.assertEqual(decode_encryption_key(encoded), self.key)
        with self.assertRaises(ValueError):
            decode_encryption_key("not base64!")
        with self.assertRaisesRegex(ValueError, "exactly 32 bytes"):
            decode_encryption_key(base64.b64encode(b"short").decode())

    def test_each_named_database_field_is_encrypted_independently(self) -> None:
        record = {field: "" for field in WORKBOOK_FIELDS}
        record["FIRST NAME"] = "Ama"
        record["LAST NAME"] = "Mensah"
        encrypted_fields = encrypt_record_fields(record, self.record_id, self.key)
        self.assertEqual(tuple(encrypted_fields), WORKBOOK_FIELDS)
        self.assertNotIn("Ama", encrypted_fields["FIRST NAME"])
        self.assertNotEqual(
            encrypted_fields["FIRST NAME"], encrypted_fields["LAST NAME"]
        )
        self.assertEqual(
            decrypt_record_fields(encrypted_fields, self.record_id, self.key), record
        )

    def test_encrypted_field_cannot_be_moved_between_columns(self) -> None:
        record = {field: "" for field in WORKBOOK_FIELDS}
        record["FIRST NAME"] = "Ama"
        encrypted_fields = encrypt_record_fields(record, self.record_id, self.key)
        encrypted_fields["LAST NAME"] = encrypted_fields["FIRST NAME"]
        with self.assertRaisesRegex(ValueError, "LAST NAME.*authenticated"):
            decrypt_record_fields(encrypted_fields, self.record_id, self.key)

    def test_encrypted_fields_require_exact_workbook_schema(self) -> None:
        with self.assertRaisesRegex(ValueError, "exactly the expected fields"):
            encrypt_record_fields({"FIRST NAME": "Ama"}, self.record_id, self.key)


class SearchTests(unittest.TestCase):
    def test_search_matches_casefolded_text_and_all_fields(self) -> None:
        records = [
            {"FIRST NAME": "Ama", "HOMETOWN": "Koforidua"},
            {"FIRST NAME": "Kwame", "HOMETOWN": "Accra"},
        ]
        self.assertEqual(search_records(records, "KOFORIDUA"), [records[0]])
        self.assertEqual(search_records(records, " kwame "), [records[1]])
        self.assertEqual(search_records(records, ""), records)


class BulkImportTests(unittest.TestCase):
    def test_exact_existing_and_repeated_candidate_records_are_skipped(self) -> None:
        existing = {field: "" for field in WORKBOOK_FIELDS}
        existing["FIRST NAME"] = "Ama"
        existing["LAST NAME"] = "Mensah"
        existing["_record_id"] = str(uuid4())
        repeated = {field: "" for field in WORKBOOK_FIELDS}
        repeated["FIRST NAME"] = "Kojo"
        repeated["LAST NAME"] = "Adu"
        similar = dict(repeated)
        similar["HOMETOWN"] = "Accra"
        other = {field: "" for field in WORKBOOK_FIELDS}
        other["FIRST NAME"] = "Yaa"

        imported = new_membership_records(
            [existing],
            [
                {field: existing[field] for field in WORKBOOK_FIELDS},
                repeated,
                dict(repeated),
                similar,
                other,
            ],
        )

        self.assertEqual(imported, [repeated, similar, other])


class WorkbookTests(unittest.TestCase):
    def test_workbook_import_normalizes_blank_header_and_skips_blank_rows(self) -> None:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.append([None, *WORKBOOK_FIELDS])
        worksheet.append([1, "Ama", "Mensah", *([None] * (len(WORKBOOK_FIELDS) - 2))])
        worksheet.append([None] * (len(WORKBOOK_FIELDS) + 1))
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "membership.xlsx"
            workbook.save(path)
            records = workbook_records(path)
        self.assertEqual(len(records), 1)
        self.assertEqual(len(records[0]), 33)
        self.assertNotIn("SOURCE COLUMN 1", records[0])
        self.assertEqual(records[0]["FIRST NAME"], "Ama")
        self.assertEqual(records[0]["LAST NAME"], "Mensah")
        self.assertEqual(records[0]["E-MAIL"], "")

    def test_workbook_rejects_unexpected_headers(self) -> None:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.append([None, "NOT THE MEMBER WORKBOOK"])
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "invalid.xlsx"
            workbook.save(path)
            with self.assertRaisesRegex(ValueError, "blank leading column"):
                workbook_records(path)

    def test_exported_workbook_matches_import_layout_and_keeps_text_literal(self) -> None:
        record = {field: "" for field in WORKBOOK_FIELDS}
        record["FIRST NAME"] = "=not a formula"
        record["LAST NAME"] = "Mensah"
        record["DATE OF BIRTH"] = "2000-02-03"
        second_record = {field: "" for field in WORKBOOK_FIELDS}
        second_record["FIRST NAME"] = "Kojo"
        second_record["LAST NAME"] = "Adu"
        exported = membership_workbook([record, second_record])

        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "export.xlsx"
            path.write_bytes(exported)
            workbook = load_workbook(path, data_only=False)
            worksheet = workbook.active
            self.assertIsNone(worksheet["A1"].value)
            self.assertEqual(
                tuple(cell.value for cell in worksheet[1][1:]),
                WORKBOOK_FIELDS,
            )
            self.assertEqual(worksheet["B2"].value, "=not a formula")
            self.assertEqual(worksheet["B2"].data_type, "s")
            workbook.close()

            self.assertEqual(
                workbook_records(BytesIO(exported)), [record, second_record]
            )


if __name__ == "__main__":
    unittest.main()
