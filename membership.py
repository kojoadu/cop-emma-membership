"""Encryption and workbook helpers for COP EMMA membership records."""

from __future__ import annotations

import base64
import binascii
from io import BytesIO
import json
import secrets
import unicodedata
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from typing import Any, BinaryIO
from uuid import UUID

from cryptography.exceptions import InvalidTag
from cryptography.hazmat.primitives.ciphers.aead import AESGCM
from openpyxl import Workbook
from openpyxl import load_workbook

ENVELOPE_VERSION = "v1"
WORKBOOK_FIELDS = (
    "FIRST NAME",
    "LAST NAME",
    "OTHER NAMES",
    "DATE OF BIRTH",
    "HOMETOWN",
    "POSTAL ADDRESS",
    "PHONE 1",
    "PHONE 2",
    "E-MAIL",
    "PLACE OF RESIDENCE",
    "STREET NAME",
    "HOUSE NUMBER",
    "RESIDENT DIRECTION",
    "EDUCATIONAL LEVEL",
    "OCCUPATION",
    "FATHER",
    "FATHER CONTACT",
    "MOTHER",
    "MOTHER CONTACT",
    "MARITAL STATUS",
    "SPOUSE NAME",
    "SPOUSE CONTACT",
    "POSITION",
    "TRANSFERRED FROM",
    "TRANSFERRED DATE",
    "BAPTIZED BY",
    "BAPTIZED DATE",
    "PLACE BAPTIZED",
    "MINISTRY",
    "COMMUNICANT",
    "GENDER",
    "ACTIVE",
    "INACTIVE",
)


def decode_encryption_key(encoded_key: str) -> bytes:
    """Decode the standard-base64 256-bit AES key used by AES-GCM."""
    if not isinstance(encoded_key, str) or not encoded_key.strip():
        raise ValueError("The record encryption key is missing.")
    try:
        key = base64.b64decode(encoded_key.strip(), validate=True)
    except (ValueError, binascii.Error) as exc:
        raise ValueError("The record encryption key must be valid base64.") from exc
    if len(key) != 32:
        raise ValueError("The record encryption key must decode to exactly 32 bytes.")
    return key


def _aad(record_id: str) -> bytes:
    try:
        canonical_id = str(UUID(record_id))
    except (ValueError, TypeError, AttributeError) as exc:
        raise ValueError("A valid membership record ID is required.") from exc
    return f"cop-emma-membership:member:{ENVELOPE_VERSION}:{canonical_id}".encode()


def _field_aad(record_id: str, field: str) -> bytes:
    if field not in WORKBOOK_FIELDS:
        raise ValueError("A valid membership field is required.")
    return _aad(record_id) + b":" + field.encode("utf-8")


def _json_default(value: Any) -> str:
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    raise TypeError(f"Unsupported membership value type: {type(value).__name__}")


def encrypt_record(record: dict[str, Any], record_id: str, key: bytes) -> str:
    if len(key) != 32:
        raise ValueError("The encryption key must be exactly 32 bytes.")
    plaintext = json.dumps(
        record, ensure_ascii=False, separators=(",", ":"), default=_json_default
    ).encode("utf-8")
    nonce = secrets.token_bytes(12)
    ciphertext = AESGCM(key).encrypt(nonce, plaintext, _aad(record_id))
    token = base64.urlsafe_b64encode(nonce + ciphertext).decode("ascii")
    return f"{ENVELOPE_VERSION}:{token}"


def decrypt_record(envelope: str, record_id: str, key: bytes) -> dict[str, Any]:
    if len(key) != 32:
        raise ValueError("The encryption key must be exactly 32 bytes.")
    aad = _aad(record_id)
    if not isinstance(envelope, str):
        raise ValueError("Encrypted membership record is malformed.")
    try:
        version, token = envelope.split(":", 1)
        payload = base64.b64decode(token.encode("ascii"), altchars=b"-_", validate=True)
    except (ValueError, binascii.Error, UnicodeEncodeError) as exc:
        raise ValueError("Encrypted membership record is malformed.") from exc
    if version != ENVELOPE_VERSION:
        raise ValueError("Unsupported membership encryption format.")
    if len(payload) < 12 + 16:
        raise ValueError("Encrypted membership record is incomplete.")
    try:
        plaintext = AESGCM(key).decrypt(payload[:12], payload[12:], aad)
    except InvalidTag as exc:
        raise ValueError(
            "Membership data could not be authenticated. Check the configured key."
        ) from exc
    try:
        record = json.loads(plaintext.decode("utf-8"))
    except (UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise ValueError("Encrypted membership record is malformed.") from exc
    if not isinstance(record, dict) or not all(
        isinstance(field, str) for field in record
    ):
        raise ValueError("Decrypted membership record has an invalid shape.")
    return record


def encrypt_record_fields(
    record: dict[str, Any], record_id: str, key: bytes
) -> dict[str, str]:
    if set(record) != set(WORKBOOK_FIELDS):
        raise ValueError("Membership records must contain exactly the expected fields.")
    if len(key) != 32:
        raise ValueError("The encryption key must be exactly 32 bytes.")
    fields: dict[str, str] = {}
    aes_gcm = AESGCM(key)
    for field in WORKBOOK_FIELDS:
        value = record[field]
        if not isinstance(value, str):
            raise ValueError(f"Membership field {field!r} must contain text.")
        nonce = secrets.token_bytes(12)
        ciphertext = aes_gcm.encrypt(
            nonce,
            value.encode("utf-8"),
            _field_aad(record_id, field),
        )
        token = base64.urlsafe_b64encode(nonce + ciphertext).decode("ascii")
        fields[field] = f"{ENVELOPE_VERSION}:{token}"
    return fields


def decrypt_record_fields(
    encrypted_fields: dict[str, Any], record_id: str, key: bytes
) -> dict[str, str]:
    if set(encrypted_fields) != set(WORKBOOK_FIELDS):
        raise ValueError("Encrypted membership fields do not match the expected layout.")
    if len(key) != 32:
        raise ValueError("The encryption key must be exactly 32 bytes.")
    aes_gcm = AESGCM(key)
    record: dict[str, str] = {}
    for field in WORKBOOK_FIELDS:
        envelope = encrypted_fields[field]
        if not isinstance(envelope, str):
            raise ValueError(f"Encrypted membership field {field!r} is malformed.")
        try:
            version, token = envelope.split(":", 1)
            payload = base64.b64decode(
                token.encode("ascii"), altchars=b"-_", validate=True
            )
        except (ValueError, binascii.Error, UnicodeEncodeError) as exc:
            raise ValueError(
                f"Encrypted membership field {field!r} is malformed."
            ) from exc
        if version != ENVELOPE_VERSION or len(payload) < 12 + 16:
            raise ValueError(f"Encrypted membership field {field!r} is malformed.")
        try:
            plaintext = aes_gcm.decrypt(
                payload[:12],
                payload[12:],
                _field_aad(record_id, field),
            )
            record[field] = plaintext.decode("utf-8")
        except (InvalidTag, UnicodeDecodeError) as exc:
            raise ValueError(
                f"Encrypted membership field {field!r} could not be authenticated."
            ) from exc
    return record


def search_records(
    records: list[dict[str, Any]], query: str
) -> list[dict[str, Any]]:
    needle = _normalize(query)
    if not needle:
        return records
    return [
        record
        for record in records
        if any(needle in _normalize(value) for value in record.values())
    ]


def new_membership_records(
    existing_records: list[dict[str, Any]],
    candidate_records: list[dict[str, Any]],
) -> list[dict[str, str]]:
    """Return unique candidate rows not already present, comparing all member fields."""

    def record_key(record: dict[str, Any]) -> tuple[str, ...]:
        return tuple(
            "" if record.get(field) is None else str(record.get(field, ""))
            for field in WORKBOOK_FIELDS
        )

    known_records = {record_key(record) for record in existing_records}
    new_records: list[dict[str, str]] = []
    for candidate in candidate_records:
        key = record_key(candidate)
        if key in known_records:
            continue
        known_records.add(key)
        new_records.append(
            {
                field: value
                for field, value in zip(WORKBOOK_FIELDS, key, strict=True)
            }
        )
    return new_records


def _normalize(value: Any) -> str:
    text = unicodedata.normalize("NFKC", str(value)).casefold()
    return " ".join(text.split())


def workbook_records(path: str | Path | BinaryIO) -> list[dict[str, str]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if not workbook.worksheets:
            raise ValueError("The workbook does not contain a worksheet.")
        worksheet = workbook.worksheets[0]
        row_iterator = worksheet.iter_rows(values_only=True)
        headers = next(row_iterator, None)
        if headers is None:
            raise ValueError("The workbook is empty.")
        normalized_headers = tuple(
            str(value).strip() if value is not None else ""
            for value in headers
        )
        if (
            len(normalized_headers) != len(WORKBOOK_FIELDS) + 1
            or normalized_headers[0] != ""
        ):
            raise ValueError(
                "The workbook must have one blank leading column followed by the "
                "expected member-detail fields."
            )
        if normalized_headers[1:] != WORKBOOK_FIELDS:
            raise ValueError(
                "The workbook headers do not match the expected membership layout."
            )

        records: list[dict[str, str]] = []
        for row in row_iterator:
            values = list(row[1 : len(WORKBOOK_FIELDS) + 1])
            if not any(value is not None and str(value).strip() for value in values):
                continue
            values.extend([None] * (len(WORKBOOK_FIELDS) - len(values)))
            record = {
                field: value.isoformat() if isinstance(value, (date, datetime))
                else "" if value is None
                else str(value).strip()
                for field, value in zip(WORKBOOK_FIELDS, values, strict=True)
            }
            records.append(record)
        return records
    finally:
        workbook.close()


def membership_workbook(records: list[dict[str, Any]]) -> bytes:
    """Serialize all member records using the importer's expected workbook layout."""
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Membership"
    worksheet.append(["", *WORKBOOK_FIELDS])
    for record in records:
        worksheet.append(["", *("" for _ in WORKBOOK_FIELDS)])
        row_number = worksheet.max_row
        for column, field in enumerate(WORKBOOK_FIELDS, start=2):
            value = record.get(field, "")
            cell = worksheet.cell(row=row_number, column=column)
            cell.value = "" if value is None else str(value)
            cell.data_type = "s"
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()
